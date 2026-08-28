import os
from datetime import datetime

import sys as _sys
_PROJECT_ROOT   = (os.path.dirname(os.path.abspath(_sys.executable))
                   if getattr(_sys, "frozen", False)
                   else os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
_USERGUIDES_DIR = os.path.join(_PROJECT_ROOT, "data", "user_guides")
_SCREENSHOTS_DIR = os.path.join(_PROJECT_ROOT, "data", "screenshots")

from core.core import log, NavigationResult
from usecases.testcase_reader import load_testcases
from usecases.testcase_runner import NavigationTester, _expand_matrix, _make_progress_cb


def generate_user_guide(app, tc_names: list[str], options: dict, progress_win=None) -> str:
    """Runs the given testcase(s) in a visible browser, screenshots every
    step, and builds a Word document (one heading/caption + screenshot
    per step).

    Returns the path to the generated .docx.
    """
    from adapters.browser.driver import start_browser, quit_browser
    from adapters.browser.login import login
    from core.core import dismiss_cookie_banner

    driver = None
    try:
        all_items = []
        url = ""
        browser_name = "chrome"
        username = ""
        password = ""
        private = False
        merged_matrix: dict = {}

        for name in tc_names:
            items, meta = load_testcases(name)
            all_items.extend(items)
            merged_matrix.update(meta.get("matrix") or {})
            if meta.get("url"):      url          = meta["url"]
            if meta.get("browser"):  browser_name = meta["browser"].strip().lower()
            if meta.get("username"): username     = str(meta["username"]).strip()
            if meta.get("password"): password     = str(meta["password"]).strip()
            if meta.get("private"):  private      = bool(meta["private"])

        if not url:
            raise ValueError("No URL found in testcase(s).")
        if browser_name not in ("chrome", "edge", "firefox"):
            browser_name = "chrome"

        merged_matrix.pop("parallel", None)
        iter_vars, _ = _expand_matrix(merged_matrix)[0]

        driver = start_browser(headless=False, browser=browser_name, private=private)
        login(driver, url, username, password)
        dismiss_cookie_banner(driver)

        progress_cb = _make_progress_cb(progress_win, ", ".join(tc_names))

        results = NavigationTester(
            driver=driver, items=all_items,
            screenshot_every_step=bool(options.get("screenshot", True)),
            screenshot_dir=_SCREENSHOTS_DIR,
            vars_context=iter_vars,
            stop_on_error=False,
            progress_callback=progress_cb).test_all()

        os.makedirs(_USERGUIDES_DIR, exist_ok=True)
        ts    = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe  = "_".join(n.replace(" ", "_") for n in tc_names)[:60]
        path  = os.path.join(_USERGUIDES_DIR, f"{safe} - {ts}.docx")
        _build_docx(path, ", ".join(tc_names), results, options)
        log.info(f"User guide created: {path}")
        return path
    finally:
        quit_browser(driver)


def _build_docx(path: str, title: str, results: list[NavigationResult], options: dict) -> None:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    want_screenshot = bool(options.get("screenshot", True))
    want_title      = bool(options.get("title", True))
    want_caption    = bool(options.get("caption", False))

    doc = Document()
    doc.add_heading(title, level=1)

    for idx, result in enumerate(results, 1):
        description = (result.description or result.method or "").strip()

        if want_title:
            doc.add_heading(f"{idx}. {description}" if description else str(idx), level=2)

        has_image = want_screenshot and result.screenshot_path and os.path.isfile(result.screenshot_path)
        if has_image:
            doc.add_picture(result.screenshot_path, width=Inches(6.0))

        if want_caption and description:
            p = doc.add_paragraph()
            if has_image:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(description)
            run.italic = True
            run.font.size = Pt(9)

    doc.save(path)
